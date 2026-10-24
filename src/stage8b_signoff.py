"""Stage 8b: outputs/kickoff_signoff.xlsx for TJ's sign-off, and reports/STAGE_8b.md.
The author columns (author_approved, author_depth, author_note) are left empty: they are for TJ only (rule 6)."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

t = pd.read_csv("outputs/kickoff_table.csv")
nz = lambda x: None if pd.isna(x) else (int(x) if float(x).is_integer() else float(x))

# ---------------- workbook
wb = Workbook(); ws = wb.active; ws.title = "signoff"
hdr = ["sidetrack", "parent", "ddr_kickoff_depth_m", "evidence_quote", "evidence_report", "method", "rule_tier", "confidence", "nod_wlbKickOffPoint_m", "nod_minus_ddr_m",
       "v2_depth_m", "alternative_depths", "note", "author_approved", "author_depth", "author_note"]
ws.append(hdr)
for r in t.itertuples():
    ws.append([r.sidetrack, r.parent, nz(r.ddr_kickoff_depth_m), r.evidence_quote, f"{r.evidence_wellbore} {r.evidence_date}", r.method, r.rule_tier, r.confidence,
               nz(r.nod_wlbKickOffPoint_m), nz(r.nod_minus_ddr_m), nz(r.v2_depth_m), None if pd.isna(r.alternative_depths) else r.alternative_depths,
               None if pd.isna(r.note) else r.note, None, None, None])
for c in ws[1]:
    c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDDDDD")
for col, w in zip("ABCDEFGHIJKLMNOP", [14, 24, 12, 62, 26, 30, 8, 11, 12, 11, 10, 36, 46, 12, 12, 40]):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in row[13:]: c.fill = PatternFill("solid", fgColor="FFF9D6")
dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"N2:N{len(t) + 1}")
ws.freeze_panes = "C2"
info = wb.create_sheet("info")
for line in ["Kick-off sign-off, one row per sidetrack (15). Fill author_approved (Y/N), author_depth (if different) and author_note; the manuscript will use only author-approved depths.",
             "Rule: the depth at which the narrative first reports new formation or departure from the parent hole; where none is stated, the depth at which the kick-off or window is started; where the narrative is empty, the measured depth at the parent-to-child transition.",
             "rule_tier: T1 = depth stated with new-formation or departure language; T2 = kick-off or window start (window top for windows); T3 = narrative empty (depth sequence).",
             "method: 'depth in text' | 'event in text, depth from md' | 'depth sequence (narrative empty)'.",
             "alternative_depths: other readings of the rule that the evidence allows; they are not applied.",
             "Evidence quotes are 15 words or fewer (cut from the raw narrative, so they can start mid-sentence). NOD = wlbKickOffPoint, snapshot 2026-10-24. Depths in m MD."]:
    info.append([line])
info.column_dimensions["A"].width = 150
wb.save("outputs/kickoff_signoff.xlsx")
chk = pd.read_excel("outputs/kickoff_signoff.xlsx", sheet_name="signoff")
print(chk.shape, "author columns filled:", int(chk[["author_approved", "author_depth", "author_note"]].notna().sum().sum()))

# ---------------- report
f = lambda x: "" if pd.isna(x) else (f"{x:,.1f}" if float(x) % 1 else f"{x:,.0f}")
fd = lambda x: "" if pd.isna(x) else (f"{x:+,.1f}" if float(x) % 1 else f"{x:+,.0f}")
rows = [f"| {r.sidetrack} | {r.parent} | {f(r.ddr_kickoff_depth_m)} | `{r.evidence_quote}` ({r.evidence_wellbore[3:]}, {r.evidence_date}) | {r.method} | {r.rule_tier} | {r.confidence} | {f(r.nod_wlbKickOffPoint_m)} | {fd(r.nod_minus_ddr_m)} |" for r in t.itertuples()]
flag = t[t.v2_status.str.startswith("NOT")]
frows = [f"| {r.sidetrack} | {f(r.v2_depth_m)} | {f(r.ddr_kickoff_depth_m)} | {r.v2_status.replace('NOT A KICK-OFF: ', '')} |" for r in flag.itertuples()]
alt = t[t.alternative_depths.notna()]
arows = [f"| {r.sidetrack} | {f(r.ddr_kickoff_depth_m)} | {r.alternative_depths} |" for r in alt.itertuples()]
vc = t.method.value_counts()
txt = f"""# Stage 8b: Kick-off under the uniform rule

Date: 2026-10-24. Branch: `rerun-2026-10`. Scripts: `src/stage8b_kickoff_table.py`, `src/stage8b_signoff.py`. Outputs: `outputs/kickoff_table.csv`, `outputs/kickoff_signoff.xlsx` (for TJ's sign-off, author columns empty). Decisions: `notes/decisions.md` items 27 to 36.

**Status: the table is final pending TJ's row-by-row sign-off. The manuscript will use only author-approved depths.**

## The rule (TJ decision 32), applied uniformly to all 15

"The depth at which the narrative first reports new formation or departure from the parent hole; where none is stated, the depth at which the kick-off or window is started; where the narrative is empty, the measured depth at the parent-to-child transition (DEPTH_SEQUENCE)."

Operationally: tier T1 = a depth is stated together with new-formation or departure language (for example "50% formation obtained at 1364 m", "kicked off at 3185m"); tier T2 = the narrative names a kick-off, window or oriented-drilling start but gives no formation or departure depth, so the start depth is used (window top for windows); tier T3 = empty narrative, measured depth at the parent-to-child transition. Evidence is read from the parent's last and the sidetrack's first reports. I chose the evidence report and depth by reading the narratives; the script checks the depth occurs in that report (text or md field) and cuts the quote.

Method wording (decision 33): "depth in text", "event in text, depth from md" (the narrative names the kick-off event but not the depth, the depth is the md of the same report), "depth sequence (narrative empty)". Counts: depth in text {vc.get('depth in text', 0)}, event in text with depth from md {vc.get('event in text, depth from md', 0)} (F-1 A, F-1 B, F-15 D), depth sequence {vc.get('depth sequence (narrative empty)', 0)} (19 ST2); {vc.sum()} in total. Tiers: T1 {int((t.rule_tier == 'T1').sum())}, T2 {int((t.rule_tier == 'T2').sum())}, T3 {int((t.rule_tier == 'T3').sum())}.

## Table (15 sidetracks, depths in m MD, NOD = wlbKickOffPoint, snapshot 2026-10-24)

| Sidetrack | Parent | DDR kick-off | Evidence quote (report, date) | Method | Tier | Confidence | NOD | NOD - DDR |
|---|---|---|---|---|---|---|---|---|
{chr(10).join(rows)}

## Rows whose depth changed under the rule

Compared with the previous version of this table (the departure-from-parent version): **one row changes. 15/9-F-11 T2: 257 m to 246 m.** The narrative of 2013-03-23 reports oriented drilling out of the cement starting at 246 m (top of cement 244 m), and gives no formation or departure depth, so tier T2 applies; 257 m was the hand-over depth of the next day. **15/9-F-1 C stays at 1,364 m** (T1: 50% formation obtained at 1,364 m). The other 13 rows are unchanged. Compared with v2, the rows that differ are the five listed below plus F-11 T2 (v2 257 m) and the new 19 A.

## v2 depths that were not kick-offs

| Sidetrack | v2 depth (m) | Re-derived (m) | What the v2 depth is |
|---|---|---|---|
{chr(10).join(frows)}

F-1 B, F-1 C and F-11 B are section starts; F-1 A and F-9 A are drilled depths of the first report. 15/9-F-11 T2 differs from v2 (257 to 246 m) because of the rule, not because v2 used a section start.

## Where the evidence allows another reading (not applied; for sign-off)

| Sidetrack | Depth used | Other depth(s) the evidence supports |
|---|---|---|
{chr(10).join(arows)}

The two window cases (19 A and F-9 A) are the main judgement: I read "where the kick-off or window is started" as the window top; reading "departure from the parent hole" as the window bottom would give 2,183 m and 428.2 m. NOD is 2,178 m and 428 m.

## 15/9-19 S (reported only, as instructed)

Structured `mdKickoff` for 19 S: 562 m in all 49 DDRs. NOD `wlbKickOffPoint` for 19 S: blank. No interpretation.

## Cross-paper correction (replaces the note of decision 28, per TJ)

"F-1 family shared intervals: all four of F-1, F-1 A, F-1 B, F-1 C share the hole above 1,364 m MD; F-1, F-1 A and F-1 B share it above approximately 2,617 to 2,620 m MD. Correct the former 'F-1 / F-1 C share formations above 3,059 m' statement in Papers 1, 3 and the barrier paper."

## Housekeeping

- `inputs/paper2_rebuild_prompt_pack_v2.md` copied from the Downloads folder (SHA-256 `18dbd3265cf9796ee26c93feb4a4486bddb1a1b392faeea4b73eddeb582884e0`; source and copy hashes identical; no secrets found by scan).
- All evidence quotes in `outputs/kickoff_table.csv` and the workbook are 15 words or fewer.

## Next

TJ reviews `outputs/kickoff_signoff.xlsx`. Stage 9 plan is in the chat reply; computing waits for TJ's go.
"""
open("reports/STAGE_8b.md", "w", encoding="utf-8").write(txt)
print("report written", len(txt))
