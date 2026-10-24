"""Stage 12 check: every number in manuscript/paper2_draft.docx must match a registry value (sign-insensitive, also as a proportion of a percentage), and no em dash may appear.
Scope: all body paragraphs and tables, excluding the Abbreviations list and the References list. Writes outputs/manuscript_registry_check.csv and prints the mismatches."""
import json
import re
import sys

import docx
import pandas as pd

E = json.load(open("outputs/numbers_registry.json", encoding="utf-8"))


def flat(v, out):
    if isinstance(v, bool) or v is None: return
    if isinstance(v, (int, float)): out.add(round(abs(float(v)), 4)); return
    if isinstance(v, str):
        for t in re.findall(r"\d[\d,]*\.?\d*", v):
            try: out.add(round(float(t.replace(",", "")), 4))
            except ValueError: pass
    elif isinstance(v, dict):
        for a, b in v.items(): flat(a, out); flat(b, out)
    elif isinstance(v, list):
        for x in v: flat(x, out)


vals = set()
for e in E:
    flat(e["value"], vals); flat(e["v2_value"], vals); flat(e["id"], vals)
vals |= {round(x / 100, 4) for x in list(vals) if 0 < x <= 100}
vals |= {round(x * 100, 4) for x in list(vals) if 0 < x < 1}                 # proportions written as percentages

path = sys.argv[1] if len(sys.argv) > 1 else "manuscript/paper2_draft.docx"
d = docx.Document(path)
paras = d.paragraphs
heads = {i: p.text.strip() for i, p in enumerate(paras) if p.style.name.startswith("Heading 1")}
skip = set()
cur_skip = False
for i, p in enumerate(paras):
    if p.style.name.startswith("Heading 1"): cur_skip = p.text.strip() in ("Abbreviations", "References")
    if cur_skip: skip.add(i)

rows = []
TOK = re.compile(r"(?<![\w.])(\d[\d,]*\.?\d*)(%?)")


def scan(text, where):
    for m in TOK.finditer(text):
        raw = m.group(1).rstrip(".,")
        try: x = round(float(raw.replace(",", "")), 4)
        except ValueError: continue
        ctx = " ".join(text[max(0, m.start() - 50): m.end() + 30].split())
        kind = ""
        if re.fullmatch(r"(19|20)\d\d", raw): kind = "year"
        elif x < 20 and "." not in raw and not m.group(2): kind = "small integer"
        elif re.fullmatch(r"\d\.\d", raw) and re.search(r"(Section|Sections)\s*$", text[:m.start()]) or (re.fullmatch(r"\d\.\d", raw) and len(text) < 80 and where.startswith("p")): kind = "section number"
        elif raw in ("25544", "1392", "985185"): kind = "identifier"
        elif raw == "701": kind = "literature figure"
        elif raw == "999.99": kind = "quoted placeholder string"
        elif raw == "95" and m.group(2): kind = "confidence level"
        rows.append(dict(where=where, token=m.group(0), value=x, status="in registry" if x in vals else (f"ignored: {kind}" if kind else "NOT IN REGISTRY"), context=ctx))


for i, p in enumerate(paras):
    if i not in skip: scan(p.text, f"p{i}")
for ti, t in enumerate(d.tables):
    for ri, row in enumerate(t.rows):
        for ci, c in enumerate(row.cells): scan(c.text, f"table {ti + 1} r{ri} c{ci}")
c = pd.DataFrame(rows); c.to_csv("outputs/manuscript_registry_check.csv", index=False)
bad = c[c.status == "NOT IN REGISTRY"]
em = sum(p.text.count("—") for p in paras) + sum(cell.text.count("—") for t in d.tables for row in t.rows for cell in row.cells)
print("tokens:", len(c), c.status.str.split(":").str[0].value_counts().to_dict(), "| em dashes:", em)
pd.set_option("display.width", 220, "display.max_colwidth", 110)
print(bad[["where", "token", "context"]].to_string(index=False))
