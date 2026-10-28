"""Stage 12 (DRAFT): build manuscript/paper2_draft.docx from inputs/paper2_v2.docx, every number from outputs/numbers_registry.json, and manuscript/CHANGES.md.
Yellow highlight = an open [TJ:] note only (the kick-off values are author-approved and carry no highlighting). Marker characters in the text below
(U+27E6 and U+27E7) delimit text to be highlighted; they are never written to the document."""
import copy
import json
import os
import re
from xml.sax.saxutils import escape

import docx
import pandas as pd
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from docx.text.run import Run

REG = {e["id"]: e for e in json.load(open("outputs/numbers_registry.json", encoding="utf-8"))}
doc = docx.Document("inputs/paper2_v2.docx")
P = list(doc.paragraphs)
L, R = "⟦", "⟧"
T, LOG, RESOLVED = {}, [], []

# section label of every paragraph (for CHANGES.md)
SEC, cur_sec = {}, "front matter"
for i, p in enumerate(P):
    if p.style.name.startswith("Heading"): cur_sec = p.text.strip()
    SEC[i] = cur_sec


def cur(i):
    return T.get(i, P[i].text)


def strip_marks(s):
    return s.replace(L, "").replace(R, "")


def log(i, old, new, rid, reason):
    assert rid is None or all(r in REG for r in rid.split(";")), rid
    LOG.append(dict(location=f"{SEC[i]} (paragraph {i})" if isinstance(i, int) else i, old=old, new=strip_marks(new), id=rid or "", reason=reason))


def sub(i, old, new, rid, reason):
    t = cur(i); assert t.count(old) == 1, (i, old, t.count(old))
    T[i] = t.replace(old, new); log(i, old, new, rid, reason)


def setp(i, new, old_desc, rid, reason):
    T[i] = new; log(i, old_desc, new if len(new) < 400 else new[:400] + " ...", rid, reason)


def dropnote(i):
    t = cur(i); k = t.index("[TJ:"); note = t[k:]
    RESOLVED.append(dict(location=f"{SEC[i]} (paragraph {i})", note=note)); T[i] = t[:k].rstrip()


def write_par(p, text):
    runs = p.runs; tmpl = copy.deepcopy(runs[0]._r)
    for r in runs: p._p.remove(r._r)
    for part in re.split("(⟦.*?⟧)", text):
        if not part: continue
        hl = part.startswith(L)
        r = copy.deepcopy(tmpl); p._p.append(r); run = Run(r, p)
        run.text = part[1:-1] if hl else part
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW if hl else None


def new_par_after(anchor_el, template_p, text, renum=True):
    text = POST(text) if renum else text
    el = copy.deepcopy(template_p._p); anchor_el.addnext(el)
    par = Paragraph(el, template_p._parent); write_par(par, text); return par


def run_sub(p, idx, old, new):
    r = p.runs[idx]; assert old in r.text, (old, r.text); r.text = r.text.replace(old, new)


def val(k):
    return REG[k]["value"]


DEL = set()                    # paragraphs deleted at flush time


def delp(i, reason):
    log(i, cur(i)[:160] + (" ..." if len(cur(i)) > 160 else ""), "(paragraph deleted)", None, reason)
    DEL.add(i)


def RENUM(t):                      # two tables were added in Section 3, so the old Tables 1 to 6 become 3 to 8
    return re.sub(r"\bTable ([1-6])\b", lambda m: f"Table {int(m.group(1)) + 2}", t)


ADJ = [("adjudications", "labels"), ("adjudication", "labelling"), ("adjudicator", "labeller"), ("adjudicated", "labelled")]


FRACG = {"1/2": "½", "1/4": "¼", "3/4": "¾", "3/8": "⅜", "5/8": "⅝"}


def POST(t):
    t = RENUM(t)
    t = re.sub(r"(\d+) (\d/\d) in\.", lambda m: m.group(1) + FRACG[m.group(2)] + " in.", t)      # hole sizes in running text: 17 1/2 in. becomes 17½ in.
    for a, b in ADJ: t = t.replace(a, b)
    return t


def subany(old, new, rid, reason):
    hits = [i for i in range(len(P)) if i not in DEL and old in cur(i)]
    assert len(hits) == 1, (old[:60], hits)
    sub(hits[0], old, new, rid, reason)


def app(i, text, rid, reason):
    old = cur(i); T[i] = old.rstrip() + " " + text; log(i, old[-80:], text, rid, reason)


# =========================================================================== front matter and abstract
sub(7, 'The archive also stores 73 narratives as the literal string "None" and 331 depths as a sentinel value, neither of which a simple null test detects.',
    'The archive also stores placeholder values in 515 fields (331 depth sentinels, 7 zero depths, 176 literal "None" strings and one date sentinel) that a simple null test does not detect.',
    "placeholders.total_cells_all_fields", "full placeholder register (decision 19); v2 counted 411 (331 + 7 + 73)")
sub(8, "The output is an open depth-time-phase dataset of 1,421 records, published with the classifier code. It supports days-versus-depth benchmarking and records-based screening of legacy wellbores, and the method transfers to other WITSML 1.4.1.1 archives.",
    f"Sidetrack kick-off depths were recovered for all 15 sidetracks by one rule; where NOD records a kick-off point (12 of 15), the recovered depth agrees within 10 m for 10. The output is an open depth-time-phase dataset of 1,421 records, with per-wellbore durations to total depth, published with the classifier code. It supports days-versus-depth benchmarking and records-based screening of legacy wellbores, and the method transfers to other WITSML 1.4-series drillReport archives.",
    "kickoff.nod_comparison;kickoff.n_sidetracks;data.witsml_schema_version", "new kick-off and duration results; abstract clause on NOD agreement replaces the F-1 family clause (TJ review); WITSML version from the files")
sub(7, "one wellbore-identity conflation", "one NOD entry date that marks a later drilling campaign", "spud.identity", "category renamed from Wellbore-identity conflation (TJ review; F-1 verified: first 2013 report 2013-07-22, no NOD pre-drill date)")
sub(10, "(Energistics 2012)", "(Energistics n.d.)", None, "Energistics reference now cites the schema version declared in the files with an access date (decision 45)")
dropnote(14)
sub(20, "dataset's WITSML 1.4.1.1 XML files into a single managed table holding 1,759 rows", "dataset's WITSML 1.4.0.0 XML files (1,759 files, 21,315,879 bytes, each recorded with a SHA-256 hash) into a single table holding 1,759 rows",
    "data.witsml_schema_version;data.n_xml_files;data.xml_total_bytes", "the files declare version 1.4.0.0 (decision 15); file count and bytes from the raw download")
sub(20, "The raw XML archive was not accessible from the compute environment during the final analysis, so the extracted table is the authoritative dataset for this study; it is published in the repository.",
    "Every field is kept as a raw string at extraction, so placeholder values stay visible, and the table is regenerated from the raw files by the published code.",
    "data.n_xml_files", "v2 said the raw XML was not accessible; it was downloaded and hashed in the rerun")
dropnote(25)
sub(25, "Spud/Conductor, Surface casing, Intermediate drilling, Intermediate casing, Reservoir drilling, Liner/Casing run", "Spud/Conductor, Surface hole and casing, Intermediate drilling, Intermediate casing, 12¼ in. and smaller hole, Liner/Casing run", None,
    "label-only change (TJ review): phase names Surface casing and Reservoir drilling renamed for display; classifier patterns and priorities are frozen and untouched")
sub(25, "because drilling reports frequently express hole and casing sizes with these glyphs.",
    "because drilling reports frequently express hole and casing sizes with these glyphs. The priority order is the order listed above. Hole and casing size separate the drilling phases: 36 in. hole and 30 in. conductor (Spud/Conductor), 26 in. hole and 20 in. casing (Surface hole and casing), 17½ in. hole (Intermediate drilling), 13⅜, 10¾ and 14 in. casing (Intermediate casing), 12¼, 8½ and 6 in. hole and coring (12¼ in. and smaller hole), and 9⅝ and 7 in. liner or casing (Liner/Casing run); the remaining phases are matched on keywords alone. Phase names describe the hole and casing sizes each pattern group matches; on Volve platform wells the 12¼ in. hole is usually the section above the reservoir. Every regular expression is listed in the repository file src/classifier_v3.py.",
    "classifier.size_patterns", "resolves the open note on priority order and hole-size rules (open item 3)")
dropnote(29)
sub(29, "was drawn from these 383 with a fixed random seed of 42.", "was drawn from these 383 with a fixed random seed of 42; the rerun of the analysis from the raw files kept this sample and its strata, after confirming that each of the 50 records maps to exactly one DDR file.",
    "audit.sample_size", "reproducibility statement for the kept sample (decision 6)")
sub(29, "Each sampled record was adjudicated as CONSTRUCTION", "Each sampled record was adjudicated by the author as CONSTRUCTION", None, "open item 4 resolved: the author made all 50 calls")
sub(31, "(4,562.1 m and 4,562.0 m).", "(4,562.1 m and 4,562.0 m). The NOD export was retrieved on 24 October 2026; its rows for the Volve wellbores are identical to an export taken on 8 September 2026 except for the register's own synchronisation date.",
    "data.nod_snapshot_date", "NOD snapshot refreshed and compared with the September snapshot")
sub(32, "new hole was kicked off from 2,911 m under", f"new hole was kicked off from 2,911 m under", "kickoff.proposed_depth.15/9-19 BT2", "kick-off depth, author-approved (unchanged value)")
sub(32, "new hole from approximately 1,536 m under", f"new hole from approximately 1,536 m under", "kickoff.proposed_depth.15/9-19 ST2", "kick-off depth, author-approved (unchanged value)")

# =========================================================================== 4 Data quality
sub(40, 'but 73 records contain the literal string "None", almost certainly the product of a Python str(None) conversion during the original extraction. A null test cannot detect a non-empty placeholder of this kind. Correcting for it, the genuine narrative population is 95.85% (1,686 of 1,759).',
    'but 73 records contain the literal string "None". The string is stored in the source XML itself (73 of the 1,759 files contain it), so it is a placeholder in the archive, not an artefact of extraction. A null test cannot detect a non-empty placeholder of this kind. Correcting for it, the genuine narrative population is 95.85% (1,686 of 1,759). The same placeholder appears in 103 forecast24hr values, so the genuine forecast population is 94.14% (1,656 of 1,759). Counting every placeholder in every extracted field gives 515 values: 331 depth sentinels, 7 zero depths, 176 literal "None" strings (73 narratives and 103 forecasts) and one date sentinel; 331 + 7 + 176 + 1 = 515.',
    "placeholders.none_sum24Hr;placeholders.none_forecast24hr;placeholders.none_total;placeholders.total_cells_all_fields;table1.forecast24hr_genuine.count;table1.forecast24hr_genuine.pct",
    "the None string is in the source archive (decision 12); forecast24hr treated like sum24Hr (decision 18); full placeholder total (decision 19)")
sub(41, "17.5 years before", "17.6 years before", "data.sentinel_years_before_first_ddr", "6,415 days / 365.25 = 17.56; TJ ruling 4 (v2 said 17.5)")
sub(41, "It is a date sentinel, not a drilling report.", "It is a date sentinel, not a drilling report. The raw files declare the ISO-8859-1 character encoding; reading them correctly repaired an encoding error that affected 103 narratives in the earlier extraction (for example the character Ø), which changes no phase.",
    "extract.old_table_encoding_rows_sum24Hr", "new finding from the raw-data rerun")

# =========================================================================== 5 Results
sub(44, "The result was checked three ways: a row-count query (1,759), a differently worded phase-distribution query returning identical per-phase counts, and manual tracing of eight randomly selected records against their matching or non-matching expressions.",
    "The ported classifier was first checked against the earlier extraction: on the narratives of the earlier table it reproduced the 1,303 classified records, every per-phase count, and the earlier phase label of all 1,759 records. Applied to narratives read directly from the raw XML it gave the same phase counts, and the 103 narratives whose encoding was repaired (Section 4) changed no phase.",
    "classifier.faithfulness_classified_on_old_text;classifier.faithfulness_row_mismatches;classifier.mojibake_phase_changes", "v2 verification steps replaced by the faithfulness check run in the rerun")
sub(52, "a direct query against the classified table returned exactly 1,421.", "recomputing the dataset from the raw XML gave the same 1,421 rows as the earlier analysis, identical in every column.", "timeline.rows", "rerun check against the v2 dataset")
sub(52, "days from first DDR, measured depth", "days from the wellbore's first report with a recorded depth, measured depth", None, "wording fix (TJ note 4): days are counted from each wellbore's first row with a usable depth")
sub(52, "Fig. 2 plots days since first DDR against measured depth for all 26 wellbores.",
    "Days are therefore counted from each wellbore's first row with a usable depth, not from its first DDR. Fig. 2 plots report number against measured depth for all 26 wellbores. The report number counts every DDR of a wellbore from its first report (1980 sentinel excluded), with or without a recorded depth, so it is the same count as the reports to TD of Table 6; only reports with a valid depth are plotted. The days_from_first_ddr column of the dataset is unchanged from the earlier analysis.", None, "Fig. 2 x-axis changed from calendar days to report count, counted as in Table 6 (TJ review)")
dropnote(61)
sub(61, "Table 4 summarises the outcome, and Fig. 4 plots the offset (WITSML minus NOD) by category.",
    f"Table 4 summarises the outcome, and Fig. 4 plots the offset by category. The offset is the WITSML spud date minus the NOD pre-drill entry date where NOD records one, since it marks the start of top-hole drilling, and otherwise the NOD main entry date. For the four wellbores whose archive contains the NOD pre-drill entry (15/9-F-4, F-5, F-14 and F-15), the first report coincides with that date (the same day, or one day later for F-4), which supports using the pre-drill date as the spud reference. The categories were assigned by inspection of the three sources for each wellbore, and the rules in the Rule column of Table 4 were formalised afterwards to describe those assignments. Applied mechanically, the rules reproduce 25 of the 26 assignments. 15/9-F-11 was assigned Exact by inspection but is unresolved under the rule: NOD records a pre-drill campaign (6 to 17 April 2012) that is absent from the DDR archive; the archive's 2013 reports describe a fresh spud from seabed, including the 30-inch conductor, so the relationship between the two campaigns cannot be determined from these records.",
    "spud.rule_thresholds_days;spud.table4;spud.unresolved;spud.predrill_first_ddr_offset_days;spud.F-11.predrill_period", "resolves the open note on spud tolerances (open item 5); NOD date rule (decision 47); F-11 stays Unresolved (decision 48); TJ review: category rules moved to the Rule column of Table 4 and the paragraph reordered (reference date, pre-drill evidence, inspection then formalised rules, 25 of 26, F-11)")
sub(64, "by three months to a full year", "by 91 to 340 days", "spud.table4", "recomputed offsets against the first report (F-12 91, F-4 92, F-7 137, F-15 340 days)")
sub(65, "an ordinary spud-to-first-report gap. NOD, however, records an entry date of 23 July 2013, nearly six years later and about four weeks before the WITSML spud date (20 August 2013) and NOD entry date (23 August 2013) of 15/9-F-1 A, the sidetrack drilled from the same slot. NOD's 15/9-F-1 record most plausibly reflects the 2013 re-entry ahead of the sidetrack campaign, not the 2007 spud. The case is therefore classified as a wellbore-identity conflation: the WITSML value is consistent with the wellbore's own chronology.",
    "an ordinary spud-to-first-report gap. The 2007 reports are top-hole only (four reports, maximum depth 226 m). F-1 was drilled to its total depth of 3,632 m in the 2013 campaign, whose first report (22 July 2013) is one day before NOD's entry date (23 July 2013). NOD records no pre-drill date for the 2007 work, unlike F-4, F-5, F-14 and F-15. NOD's date therefore marks the later campaign, not the 2007 spud, and the case is classified as NOD date marks later campaign: the WITSML value is consistent with the wellbore's own chronology.",
    "f1.reports_2007;f1.reports_2013;f1.spud_section_statement;f1.first_report_2013;f1.nod_main_entry_date;f1.nod_predrill_entry_date", "TJ ruling 3 (decision 42); TJ review: category renamed, F-1 paragraph rewritten after verifying that the first 2013 report is 2013-07-22 and that NOD records no pre-drill date for F-1; the re-entry-ahead-of-the-sidetrack reading removed")

# ---- 5.6 kick-off (rewritten), 5.7 duration (new)
dropnote(67)
setp(67, "Because the structured kick-off field is sparse, kick-off depth was recovered from the DDRs. The kick-off depth of a sidetrack is the depth at which it leaves its parent wellbore. Kick-off here means the depth at which a sidetrack departs its parent wellbore, not the depth at which an original wellbore begins to build angle; the structured kick-off field of 15/9-19 S (562 m in all its reports) is consistent with the latter and is not used. It was read from the parent's last and the sidetrack's first reports by one rule applied to all sidetracks: the depth at which the narrative first reports new formation or departure from the parent hole; where none is stated, the depth at which the kick-off or window is started (for casing exits by a milled window, the top of the window); and where the narrative is empty, the measured depth at the parent-to-child transition.",
     "v2 5.6 first paragraph (twelve free-text and two depth-sequence recoveries)", "kickoff.ruling_window_top;kickoff.19S_structured_field_m", "Section 5.6 rewritten from the kick-off table with the uniform rule (decisions 27, 32, 40); TJ review: sentence on the meaning of kick-off and the 15/9-19 S structured field added (decision 55)")
dropnote(69)
setp(69, f"NOD's wlbKickOffPoint is available for 12 of the 15 sidetracks. It equals the recovered depth for 6 (15/9-19 A, F-1 A, F-1 C, F-15 A, F-15 B and F-15 D), differs by 10 m or less for 4 (15/9-F-9 A, F-11 A, F-11 B and F-15 C), and by more for 2 (15/9-19 B, 31 m; 15/9-F-1 B, 13 m). NOD has no kick-off value for 3 (15/9-19 BT2, 19 ST2 and F-11 T2). The least certain depths are 15/9-19 ST2, which rests on the measured depth alone, and 15/9-F-11 B, where the narrative reports cement and new formation drilled from 2,531 to 2,589 m and the end of the interval is taken as the departure.",
     "v2 5.6 paragraph on NOD's wlbKickOffPoint (operational versus geometric definition)", "kickoff.nod_comparison;kickoff.F-11 B.interval_m", "the 'operational versus geometric' paragraph is deleted (decision 27); NOD comparison counts from the kick-off table")
sub(75, 'and the narrative string "None" together affect 411 field values (331 + 7 + 73).', 'and the string "None" (in the narrative and forecast fields), together with the date sentinel, affect 515 field values (331 + 7 + 176 + 1).',
    "placeholders.total_cells_all_fields", "full placeholder register (decision 19)")
sub(72, "concentrated in drilling progress reported without a hole size.", "concentrated in drilling progress reported without a hole size. Table 6 gives elapsed days, reports and metres per report-day to total depth for each wellbore as a first benchmark, with the gap flag marking wellbores whose elapsed time cannot be read as drilling time.",
    "duration.table", "new duration results")
sub(82, "such as 15/9-F-1 and 15/9-F-1 C, are not independent penetrations above their kick-off depth.", f"such as the four wellbores of the 15/9-F-1 family, which share the hole above 1,364 m MD, are not independent penetrations above their kick-off depths.",
    "kickoff.F-1_family_shared_hole_above_m", "decision 34 replaces the 'F-1 / F-1 C share formations above 3,059 m' statement")
sub(84, "any WITSML 1.4.1.1 archive", "WITSML 1.4-series drillReport archives (tested here on 1.4.0.0)", "data.witsml_schema_version", "WITSML version from the files; TJ review: tested only on 1.4.0.0")
sub(72, "Days-versus-depth curves (Fig. 2)", "Depth-versus-report-count curves (Fig. 2)", None, "Fig. 2 x-axis is the report number (TJ review)")
sub(88, "2013 re-entry reports ahead of the F-1 A sidetrack", "2013 drilling campaign reports", None, "F-1 stratum wording (TJ review)")
sub(98, "one apparent error is a wellbore-identity conflation in the register", "one apparent error is a NOD entry date that marks a later drilling campaign", "spud.identity", "category renamed (TJ review)")
sub(87, "was adjudicated by a single reviewer, and", "was adjudicated by a single reviewer, the author, and", None, "open item 4 resolved")
setp(90, f"Kick-off depths were read from the narratives by one rule, and two rest on weaker evidence: 15/9-19 ST2 on the measured depth alone and 15/9-F-11 B on an interval of cement and new formation (2,531 to 2,589 m). NOD gives no kick-off point for three sidetracks. All four wellbores of the 15/9-F-1 family (F-1, F-1 A, F-1 B and F-1 C) share the hole above 1,364 m MD, and F-1, F-1 A and F-1 B share it above approximately 2,617 to 2,620 m MD; any formation-level analysis treating them as independent must account for this. Metres per report-day are rates per reported day, not per elapsed day, and the elapsed time to total depth is affected by gaps between reports in the archive (the gap flag in Table 6).",
     "v2 limitation: two depth-sequence kick-offs; F-1 and F-1 C shared interval above the F-1 C kick-off", "kickoff.F-1_family_shared_hole_above_m;kickoff.F-1_F-1A_F-1B_shared_hole_above_m;kickoff.F-11 B.interval_m", "decisions 34 and 40; duration caveat added")
sub(92, "The analysis uses the extracted table rather than the raw XML, so any extraction error upstream of that table would propagate.", "The analysis starts from the raw XML, extracted by the published code with every field kept as a raw string, but an extraction error would still propagate to every result.",
    "data.n_xml_files", "v2 said the analysis used the extracted table; the rerun starts from the raw files")
sub(99, "affect 411 field values", "affect 515 field values", "placeholders.total_cells_all_fields", "full placeholder register (decision 19)")
sub(101, "and use the dataset for phase-duration benchmarking across the campaign.", "and extend the duration results to phase-level durations across the campaign.", None, "duration results now reported (Section 5.7)")
dropnote(104)
setp(104, f"Generative AI (Anthropic Claude, used through Claude Code) built and ran the reproducible analysis pipeline from the raw data and was used to assist with debugging, checking arithmetic and editing the manuscript text. The sidetrack kick-off depths were proposed by the AI tool from the narratives under one rule. The author rejected the tool's first proposals for three sidetracks (15/9-19 A, 15/9-F-9 A and 15/9-F-11 B) and kept the earlier values, then reviewed all 15 depths and approved every one without change. The author made all 50 audit adjudications, reviewed all AI-assisted content, and takes full responsibility for the paper.",
     "v2 AI disclosure (writing, debugging, arithmetic, editing; note about audit adjudication and figures)", None, "AI disclosure matches what happened: AI proposed, author reviewed all 15 rows of the sign-off workbook, approved all 15, changed none, marked none N (decisions 37, 40 and 52)")
dropnote(106)
setp(106, f"The classifier, the extraction and analysis code, the 1,421-row timeline dataset, the audit adjudications (file name, wellbore, date, stratum, call and miss reason), the NOD comparison tables, the kick-off and duration tables, and a numbers registry tracing every reported figure to the script that produced it are available at https://github.com/985185/volve-well-construction-timeline. DDR narrative text is not redistributed: the Volve data are available from Equinor under the Equinor Open Data Licence (https://www.equinor.com/energy/volve-data-sharing), and the repository holds only file names, dates, depths, categories and quotes of 15 words or fewer. Comparisons with the author's earlier extraction are reported but are not needed to reproduce any figure.",
     "v2 availability statement (audit sample with adjudications, extracted DDR table, NOD tables, registry)", None, "licence and no-narrative-text policy (decisions 26 and 46); TJ review: sentence on the earlier-extraction comparisons added (decision 58); the extracted table is no longer published because it holds narrative text; Stage 13 check: every item listed exists in the repository (decision 57), so the confirmation note is removed")
setp(143, f"Every quantitative figure in this paper traces to an entry in the numbers registry published in the repository ({REG['registry.n_entries']['value']} entries, each with the producing script and function, the SHA-256 of its input files and an execution timestamp). The audit estimates in Section 5.2 are derived arithmetically from registry entries, and the arithmetic is shown in the text.",
     "v2 Appendix A (59 entries, each with the exact query and execution timestamp)", "registry.n_entries", "registry rebuilt with 224 entries; resolves the open note on the attainable-rate and placeholder-total entries")
RESOLVED.append(dict(location=f"{SEC[143]} (paragraph 143)", note=P[143].text[P[143].text.index("[TJ:"):]))

# references and captions
assert P[119].runs[0].text.rstrip().endswith("2012."), P[119].runs[0].text
run_sub(P[119], 0, "2012.", "n.d.")
run_sub(P[119], 1, "Version 1.4.1.1", "data schema version 1.4.0.0, as declared in the Volve DDR files")
run_sub(P[119], 2, "witsml-data-standards/.", "witsml-data-standards/ (accessed 24 October 2026).")
log("References (paragraph 119)", "Energistics. 2012. ... Version 1.4.1.1", "Energistics. n.d. ... data schema version 1.4.0.0, as declared in the Volve DDR files ... (accessed 24 October 2026)", "data.witsml_schema_version",
    "release label and year of 1.4.0.0 could not be verified (decision 45), so the schema version is cited with an access date")
sub(123, "(accessed 8 September 2026)", "(accessed 24 October 2026)", "data.nod_snapshot_date", "snapshot refreshed")
t129 = P[129]; run_sub(t129, 0, "Coley, C.,", "Coley, C.J.,"); t129._p.remove(t129.runs[1]._r)
RESOLVED.append(dict(location="References (paragraph 129)", note="[TJ: verify] (Sidahmed et al. 2015); verified through the Crossref record of DOI 10.2118/173429-MS, decision 44"))
log("References (paragraph 129)", "Sidahmed, M., Coley, C., and Shirzadi, S. 2015 ... [TJ: verify]", "Sidahmed, M., Coley, C.J., and Shirzadi, S. 2015 ...", None, "title and authors verified through Crossref (OnePetro returned 403); citation kept")
sub(137, "Fig. 2: Days since first DDR versus measured depth for all 26 wellbores, coloured by phase.", "Fig. 2: Report number (1 = the wellbore's first report, as counted in Table 6) versus measured depth for all 26 wellbores, coloured by hole section (other phases and unclassified reports in grey).", None, "Fig. 2 x-axis is the report number (TJ review)")
sub(139, "for the 24 matched wellbores, with 1:1 reference line.", "for the 24 matched wellbores, with 1:1 reference line. The two unmatched wellbores (15/9-19 B and 15/9-F-11) are not plotted.", None, "title removed from the image; information moved to the caption (TJ review)")
sub(141, "Fig. 4: Spud date offset (WITSML minus NOD) for 25 wellbores, coloured by validation category.", "Fig. 4: Spud date offset (WITSML spud date minus the NOD pre-drill entry date where recorded, otherwise the NOD main entry date) for 25 wellbores, coloured by validation category. The bar of 15/9-F-1 (offset -2,064 days) is clipped at -400 days and labelled with its value.", "spud.table4", "one offset definition, outcome-independent NOD date rule (decisions 41 and 47)")

# =========================================================================== TJ review, round 2 (structure and wording)
sub(6, "The final classifier version assigns", "It assigns", None, "no classifier versions named (TJ review)")
sub(7, "because the archive does not cover", "because the DDR archive does not cover", None, "wording (TJ review)")

# 1 Introduction: one paragraph, contribution reduced to one sentence
t11 = cur(11); k11 = t11.index("The paper is written to be self-contained")
CONTRIB = "The contribution is the combination, on a real, openly licensed field archive, of a transparent classifier, an estimate of what that classifier misses, validation against the regulator's register for every wellbore rather than a subset, and a description of the data-quality problems that any other user of the archive will meet. "
setp(10, cur(10) + " " + t11[:k11] + CONTRIB + t11[k11:], "v2 Introduction, three paragraphs (context, scope, contribution and roadmap)", None, "TJ review: Introduction merged into one paragraph; the contribution paragraph reduced to one sentence; roadmap sentences removed")
delp(11, "merged into the first Introduction paragraph (TJ review)")
delp(12, "reduced to one sentence inside the first Introduction paragraph (TJ review)")

# 2 Background
t14 = cur(14); T[14] = "Several studies have extracted structured information from the free text of drilling reports." + t14[t14.index(" Sidahmed"):]
log(14, "Text analytics on drilling reports has developed alongside operators' interest in extracting structured information from legacy narratives.", "Several studies have extracted structured information from the free text of drilling reports.", None, "opening sentence replaced (TJ review)")
t17 = cur(17); a17 = t17.index("The present paper differs in three respects."); b17 = t17.index(" a stratified audit")
T[17] = t17[:a17] + "The present paper differs in three respects: (1) it uses a deterministic, auditable classifier, trading model sophistication for transparency of every decision; (2) it works exclusively with a real, openly licensed archive; and (3) it pairs classification with" + t17[b17:]
log(17, "three respects as three sentences", "three respects numbered (1), (2), (3) inline", None, "TJ review")

# 3 Section 3.1
sub(20, "The Volve field (block 15/9, Norwegian North Sea), operated by Equinor (formerly Statoil), was released as an open dataset after decommissioning (Equinor 2018). The DDR archive was extracted from the dataset's", "The DDR archive was extracted from the Volve dataset's", None, "sentence repeating the Volve introduction removed (TJ review)")

# 4 Section 3.6 deleted
delp(33, "Section 3.6 deleted (TJ review); it was the last subsection of Section 3, so no other subsection is renumbered")
delp(34, "Section 3.6 deleted (TJ review)")

# 6 and 7 Section 3.3
sub(25, "Before matching, vulgar-fraction characters (½, ¼, ¾, ⅛) are replaced with ASCII equivalents, because drilling reports frequently express hole and casing sizes with these glyphs.",
    "Before matching, fraction characters such as ½ and ¼, routine in the mixed units of drilling reports, are stored as single special characters that a pattern written as \"1/2\" does not match, so they are converted first.", None, "wording (TJ review)")
delp(26, "paragraph on three classifier versions deleted: classifier versions are not named in the text (TJ review)")
sub(28, "complete v3 pattern set", "complete pattern set", None, "no classifier versions named (TJ review)")

# 9 Section 3.4 in plain language
setp(29, "A stratified sample of 50 records was drawn from these 383. The unclassified reports were split into two strata, an early construction group (reports within 730 days of the wellbore's first report, 296 records) and a later production group (reports more than 730 days after it, 87 records; 296 + 87 = 383), and a random sample was drawn from each in proportion to its size: 39 construction-group and 11 production-group records (39 + 11 = 50). The random draw is recorded in the repository so the same sample can be drawn again; the rerun of the analysis from the raw files kept this sample and its groups, after confirming that each of the 50 records maps to exactly one DDR file. The author read each sampled report and labelled it construction if it described work that builds the well: drilling ahead, coring, running or cementing casing, setting plugs, fishing, or completion. Reports that only described handling the BOP or riser, tripping, waiting or production were labelled non-construction, unless they also named a hole or casing size.",
     "v2 Section 3.4, second paragraph (stratified sample, seed, adjudication rule)", "audit.sample_size;audit.stratum_construction;audit.stratum_production;audit.sample_construction;audit.sample_production;audit.strata_cutoff_days;audit.frame_size",
     "TJ review: plain-language description of the stratified sample and the labelling rule; the seed is no longer quoted; adjudicated replaced by labelled")

# 14 Section 3.2, 3.5 and references
sub(23, "(NOD 2026e)", "(NOD 2025)", "external.nod_guidelines_revision_date", "the guidelines are cited by their own revision (28 March 2025, stated in the document's revision history), not as a FactPages page (TJ review)")
app(31, "All comparisons are against the current register as accessed in October 2026.", "data.nod_snapshot_date", "TJ review")
IDX_E = [i for i in range(len(P)) if P[i].text.startswith("Norwegian Offshore Directorate. 2026e")]; assert len(IDX_E) == 1
IDX_A = [i for i in range(len(P)) if P[i].text.startswith("Norwegian Offshore Directorate. 2026a")]; assert len(IDX_A) == 1
sub(IDX_E[0], "2026e. Guidelines for Designation of Wells and Wellbores.", "2025. Guidelines for Designation of Wells and Wellbores (revision of 28 March 2025).", "external.nod_guidelines_revision_date", "TJ review: cited by its own revision date; 2026a to 2026d now denote FactPages pages only, each with an access date")

# 12 Section 4: a real before and after phrase
_ex = val("extract.encoding_example")
sub(41, "(for example the character Ø)", f"(for example, the earlier table held '{_ex['earlier_table']}' where the raw file reads '{_ex['raw_file']}')", "extract.encoding_example", "real before and after phrase (TJ review)")
sub(41, "the vulgar fraction one-half (½)", "the fraction character one-half (½)", None, "wording (TJ review)")

# 5.1 and 5.2
sub(44, "Applying v3 to all 1,759 DDRs", "Applying the classifier to all 1,759 DDRs", None, "no classifier versions named (TJ review)")
subany("Phase distribution, v3 classifier", "Phase distribution", None, "no classifier versions named (TJ review)")
OLD50 = cur(49)
setp(49, "26 of the 50 sampled reports (52%) were construction work. Scaled to the 383-report frame this is about 199 reports, so a complete pattern set would classify about 85.4% of the 1,759 reports (95% confidence interval 82.5 to 88.3%). The arithmetic, including the stratum-weighted check, is in Appendix B.",
     "v2 Section 5.2, arithmetic paragraph", "audit.construction_calls;audit.sample_size;audit.frame_size;audit.residual_miss_ddrs;audit.attainable_rate_pct;audit.attainable_ci_low_pct;audit.attainable_ci_high_pct",
     "TJ review: plain result kept in the text; full arithmetic and the stratum-weighted check moved to Appendix B")
subany("and the arithmetic is shown in the text.", "and the arithmetic is shown in Appendix B.", "registry.n_entries", "TJ review")

# 13 Section 5.4: 15/9-19 S
app(59, "The archive's reports for this wellbore are incomplete throughout, with 39 reports over 91 days (Table 6), and the cause cannot be determined from the records.", "duration.table", "TJ review: no cause is proposed")

# limitations
subany("but its effect was not separated from other pattern changes between versions.", "but its effect was not separated from other refinements of the pattern set.", None, "no classifier versions named (TJ review)")

# =========================================================================== TJ review, round 3
# 2 Background: Karnot et al. (2025)
IK = [i for i in range(len(P)) if cur(i).startswith("The nearest comparable published work is Karnot")]; assert len(IK) == 1
setp(IK[0], "Karnot et al. (2025) combine change-point detection, semi-supervised natural language processing and sequence mining for root-cause analysis of non-productive time, applied to 701 DDRs from 10 wells. Their Dataset section states that \"all analyses used synthetic datasets constructed to emulate typical drilling records\". Their sequence-mining section, which they describe as a proof-of-concept analysis, states that \"rigorous validation and quantitative performance assessment are left for future work\". Their abstract and conclusions nevertheless report an average 18.5% reduction in non-productive time in field applications.",
     "v2 Section 2, Karnot et al. paragraph", "literature.karnot_2025", "TJ review round 3 and 4: description replaced; both quotations and the 18.5% verified by the author against the source PDF (26 October 2026); the first quotation is from the Dataset section, the second from the sequence-mining section (a proof-of-concept analysis)")
subany("The common limitation across the published DDR work is that it relies on proprietary or synthetic data, so its results", "The common limitation across the published DDR work is that its data are proprietary or synthetic rather than an openly available archive, so its results", None, "kept consistent with the Karnot et al. description (TJ review round 3)")

# 3 F-12
IF = [i for i in range(len(P)) if cur(i).startswith("In the four error cases")]; assert len(IF) == 1
tF = cur(IF[0]); a0 = tF.index("In the four error cases"); a1 = tF.index("Two sidetracks, 15/9-F-11 A")
T[IF[0]] = tF[:a0] + "In the four error cases the structured spud_date contradicts the wellbore's own first report and its narrative by 91 to 340 days, while the first report and the NOD date agree with each other. 15/9-F-12 is the clearest example: the first report (14 June 2007), which describes making up the 36 in. hole opener, and the NOD entry date agree, and only the structured field (15 March 2007) disagrees. " + tF[a1:]
log(IF[0], "three sources agree on 14 June 2007 (first DDR, narrative, NOD)", "the first report and the NOD entry date agree; the first report and its narrative are one document", "spud.F-12.witsml_spud;spud.F-12.first_ddr", "TJ review round 3: the first DDR and the free-text evidence are the same report, so 'three sources' overstated the agreement")

# availability statement: the 15-word limit is for the repository's data files
subany("and the repository holds only file names, dates, depths, categories and quotes of 15 words or fewer.", "and the data files of the repository hold only file names, dates, depths, categories and quotes of 15 words or fewer.", None, "the 15-word limit applies to data files, not to quotations in the manuscript (decision 67)")

# post-processing of every original paragraph: table numbers (+2) and the labelling words
for i in range(len(P)):
    if i in DEL: continue
    t = cur(i); nt = POST(t)
    if nt != t: T[i] = nt
log("Tables", "Tables 1 to 6", "Tables 3 to 8", None, "two tables were added in Section 3 (the example report, Table 1, and the phases, Table 2), so every later table number moves up by two; every reference in the text moves with it")
log("Whole text", "adjudicated, adjudication, adjudications, adjudicator", "labelled, labelling, labels, labeller", None, "plain language (TJ review)")

# =========================================================================== final review fixes (28 October 2026)
def ix(prefix):
    hits = [i for i in range(len(P)) if i not in DEL and cur(i).startswith(prefix)]
    assert len(hits) == 1, (prefix, hits)
    return hits[0]


# 1 no earlier extraction / analysis / table, except one sentence in 5.1
i = ix("A character-level survey of sum24Hr"); t = cur(i); k = t.index("The raw files declare the ISO-8859-1")
T[i] = t[:k] + f"The raw files declare the ISO-8859-1 character encoding; reading them with the wrong encoding turns ½ into Â½ (for example '{_ex['earlier_table']}' instead of '{_ex['raw_file']}'), so the encoding must be set explicitly."
log(i, t[k:], T[i][k:], "extract.encoding_example", "final review: the encoding point rephrased as a lesson for users; the comparison with the earlier table and the count of 103 removed")
i = ix("Applying the classifier to all 1,759 DDRs"); t = cur(i); a = t.index("The ported classifier was first checked"); b = t.index("Table 4 gives")
T[i] = t[:a] + "The classifier was also checked against an independent earlier extraction of the same archive by the author, with identical results for all 1,759 reports. " + t[b:]
log(i, t[a:b], "The classifier was also checked against an independent earlier extraction of the same archive by the author, with identical results for all 1,759 reports.", "classifier.faithfulness_classified_on_old_text;classifier.faithfulness_row_mismatches", "final review: the one remaining mention of the earlier extraction")
subany("; recomputing the dataset from the raw XML gave the same 1,421 rows as the earlier analysis, identical in every column.", ".", "timeline.rows", "final review: comparison with the earlier analysis removed")
subany(" The days_from_first_ddr column of the dataset is unchanged from the earlier analysis.", "", None, "final review: sentence removed")
subany("Comparisons with the author's earlier extraction are reported but are not needed to reproduce any figure.", "The check against the author's earlier extraction (Section 5.1) is reported but is not needed to reproduce any figure.", None, "final review: matches the single mention in Section 5.1")

# 2 and 6 Section 5.2 merged; construction group and production group; plain words
i49 = ix("Of the 50 sampled records"); i50 = ix("26 of the 50 sampled reports"); i51 = ix("The 73 EMPTY_RECORD DDRs")
setp(i49, "Of the 50 sampled reports, 26 (52%) were labelled construction and 24 (48%) non-construction. The construction group (39 reports) split 23 construction to 16 non-construction, and the production group (11 reports) split 3 to 8. The 26 construction reports that the classifier missed fall into three kinds: genuine drilling reported without a hole size (9); other operations, such as hydraulic fracturing or fishing, with no corresponding pattern (10); and non-standard terms, such as \"CMT\" for cement (7); 9 + 10 + 7 = 26. Scaled to the 383-report frame, this is about 199 reports, so a complete pattern set would classify about 85.4% of the 1,759 reports (95% confidence interval 82.5 to 88.3%). The arithmetic, including the stratum-weighted check, is in Appendix B.",
     "Section 5.2, two paragraphs (counts and plain result)", "audit.construction_calls;audit.sample_size;audit.sample_construction;audit.sample_production;audit.construction_calls_construction_era;audit.construction_calls_production_era;audit.miss_reason.NO_HOLE_SIZE;audit.miss_reason.OTHER_MISS;audit.miss_reason.NONSTANDARD_TERMINOLOGY;audit.residual_miss_ddrs;audit.attainable_rate_pct;audit.attainable_ci_low_pct;audit.attainable_ci_high_pct",
     "final review: the two paragraphs merged so that 26 of 50 appears once; construction group and production group; code labels replaced by plain words")
delp(i50, "merged into the first paragraph of Section 5.2 (final review)")
sub(i51, "The 73 EMPTY_RECORD DDRs are excluded", "The 73 DDRs whose narrative is the literal string \"None\" are excluded", None, "final review: plain words instead of the code label")
subany("were labelled NON_CONSTRUCTION, so the attainable rate", "were labelled non-construction, so the attainable rate", None, "final review: plain words")
i = ix("Two records were assigned to a stratum"); t = cur(i)
T[i] = t.replace("to a stratum that does not reflect", "to a group that does not reflect").replace("production-era stratum", "production group")
log(i, "stratum, production-era stratum", "group, production group", None, "final review: construction group and production group")

# 3 Section 5.6 definition
subany("The kick-off depth of a sidetrack is the depth at which it leaves its parent wellbore. Kick-off here means", "Kick-off here means", None, "final review: one definition sentence")
subany("is consistent with the latter and is not used. It was read from the parent's last", "is consistent with the latter and is not used. The kick-off depth was read from the parent's last", None, "final review: the dangling pronoun fixed")

# 4 Section 3.3
i = ix("Each DDR's sum24Hr narrative is classified"); t = cur(i)
a = t.index("Before matching, fraction characters"); b = t.index("The priority order is the order listed above.")
t = t[:a] + "Fraction characters such as ½ and ¼, routine in the mixed units of drilling reports, are stored as single special characters that a pattern written as \"1/2\" does not match, so they are converted to that form before matching. " + t[b:]
a = t.index("Hole and casing size separate the drilling phases:"); b = t.index("the remaining phases are matched on keywords alone. ") + len("the remaining phases are matched on keywords alone. ")
T[i] = t[:a] + t[b:]
log(i, "fraction sentence; list of hole and casing sizes", "fraction sentence rewritten; the size list deleted (Table 2 carries it)", "classifier.phase_table", "final review")

# 8 Abstract
app(ix("Sidetrack kick-off depths were recovered for all 15 sidetracks by one rule; where NOD"), "The results show that neither the structured fields nor the national register alone gives a complete well history; the narrative reports are needed to correct both.", None, "final review: closing sentence of the abstract")

# reference verification, in-text (28 October 2026)
sub(ix("Several studies have extracted"), "Hoffimann et al. (2017) applied deep natural language processing to classify sentences in drilling reports from hundreds of wells in a single field as", "Hoffimann et al. (2018) applied deep natural language processing to classify sentences in 9,670 daily reports from a field of 303 wells as", "literature.hoffimann_2018", "published SPE version cited (SPE-191505-MS, 2018); description with the verified figures (author check, 28 October 2026)")
i = ix("Several studies have extracted"); t = cur(i); k = t.index("Zhang et al. (2022) compared")
T[i] = t[:k] + "Zhang et al. (2022) built a rule-based detector for five unscheduled events (stuck pipe, lost circulation, water influx, gas influx and H2S) in about 128,000 free-text remarks from daily mud reports; the rule-based model was precise but missed events, with accuracy falling to 65.18% on a larger test set, and a hybrid that used its output to train a weakly supervised deep learning model improved detection by more than 20 percentage points."
log(i, t[k:], T[i][k:], "literature.zhang_2022", "description corrected against the source PDF (author check, 28 October 2026)")
sub(ix("The Volve dataset itself"), "Tunkiel et al. (2020)", "Tunkiel et al. (2021)", None, "year corrected (J Pet Sci Eng 196: 108069, 2021; author check, 28 October 2026)")
app(ix("The common limitation across the published DDR work"), "The audit in this paper measures the same weakness of rule-based classification, the events it misses, directly on the unclassified residual.", "literature.zhang_2022", "added with the corrected Zhang et al. description")

# 10 hole sizes in running text: 17 1/2 in. becomes 17½ in. (quotations are not touched: they carry inch marks, not 'in.')
GL = {"1/2": "½", "1/4": "¼", "3/4": "¾", "3/8": "⅜", "5/8": "⅝"}
for j in range(len(P)):
    if j in DEL: continue
    tj = cur(j); nj = re.sub(r"(\d+) (\d/\d) in\.", lambda m: m.group(1) + GL[m.group(2)] + " in.", tj)
    if nj != tj: T[j] = nj; log(j, "17 1/2 in. style", "17½ in. style", None, "final review: hole sizes written with the fraction character in running text")


# flush the edited paragraphs
for i, text in T.items():
    write_par(P[i], text)
for i in sorted(DEL):
    P[i]._p.getparent().remove(P[i]._p)
P[IDX_A[0]]._p.addprevious(P[IDX_E[0]]._p)          # reference order: the 2025 guidelines come before the 2026 FactPages pages

# figures replace the placeholders
for i, fn, w in [(134, "figure1_depth_vs_time.png", 6.5), (136, "figure2_days_vs_depth.png", 6.0), (138, "figure3_td_validation.png", 5.2), (140, "figure4_spud_offset.png", 5.4)]:
    RESOLVED.append(dict(location=f"Figures (paragraph {i})", note=P[i].text))
    for r in P[i].runs: P[i]._p.remove(r._r)
    P[i].add_run().add_picture(f"outputs/figures/{fn}", width=Inches(w))
    sp = P[i]._p.pPr.find(qn("w:spacing"))
    if sp is not None and sp.get(qn("w:line")) is not None: del sp.attrib[qn("w:line")]       # a multiple line height clips inline images
    P[i].paragraph_format.keep_with_next = True                                               # the image stays with its caption below it
    P[i + 1].paragraph_format.keep_with_next = None                                           # the caption does not hold on to the next figure

# =========================================================================== tables
t1 = doc.tables[0]
def set_cell(cell, text):
    par = cell.paragraphs[0]
    if par.runs:
        par.runs[0].text = text; return
    tbl = cell._tc.getparent().getparent()                    # empty cell: copy the run format of a non-empty cell in the same column
    col = cell._tc.getparent().findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc").index(cell._tc)
    for tr in tbl.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr"):
        tcs = tr.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc")
        for r in tcs[col].iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r"):
            nr = copy.deepcopy(r); par._p.append(nr); Run(nr, par).text = text; return
    par.add_run(text)
set_cell(t1.rows[3].cells[1], "99.94%"); set_cell(t1.rows[3].cells[2], "1,758"); set_cell(t1.rows[3].cells[3], "1 record missing (1980 sentinel)")
log("Table 1, row rig", "100.0% / 1,759", "99.94% / 1,758, note: 1 record missing (1980 sentinel)", "table1.rig_npd_name_alias.count;table1.rig_npd_name_alias.pct", "the v2 value counted the constant column 'name' (decision 17)")
set_cell(t1.rows[4].cells[0], "forecast24hr (genuine)"); set_cell(t1.rows[4].cells[1], "94.14%"); set_cell(t1.rows[4].cells[2], "1,656"); set_cell(t1.rows[4].cells[3], '103 literal "None" strings excluded')
log("Table 1, row forecast24hr", "100.0% / 1,759", '94.14% / 1,656, note: 103 literal "None" strings excluded', "table1.forecast24hr_genuine.count;table1.forecast24hr_genuine.pct", "treated like sum24Hr (decision 18, deliberate change)")
for row in doc.tables[1].rows:
    c0 = row.cells[0].text.strip()
    if c0 in ("Reservoir drilling", "Surface casing"): set_cell(row.cells[0], {"Reservoir drilling": "12¼ in. and smaller hole", "Surface casing": "Surface hole and casing"}[c0])
log("Table 2", "Reservoir drilling; Surface casing", "12¼ in. and smaller hole; Surface hole and casing", None, "label-only change (TJ review): counts unchanged (256 and 66), classifier patterns frozen")
t4 = doc.tables[3]
set_cell(t4.rows[1].cells[1], "6"); set_cell(t4.rows[1].cells[2], "15/9-19 A, F-10, F-15 A, F-15 B, F-15 C, F-15 D")
set_cell(t4.rows[2].cells[2], "Administrative spud precedes operational start by days to weeks (15/9-19 B, F-1 A, F-1 B, F-1 C, F-5, F-9 A, F-11 B, F-14)")
set_cell(t4.rows[6].cells[1], "5"); set_cell(t4.rows[6].cells[2], "15/9-19 BT2, 19 S, 19 ST2, F-9, F-11 (NOD records a pre-drill campaign, 6 to 17 April 2012, that is absent from the DDR archive; the 2013 reports describe a fresh spud from seabed, so the relationship between the campaigns cannot be determined from these records)")
log("Table 4, row Unresolved", "4, 15/9-19 series", "5: 15/9-19 BT2, 19 S, 19 ST2, F-9, F-11", "spud.unresolved", "correction of the v2 detail cell (19 B is a definitional offset, F-9 is unresolved) and 15/9-F-11 moves from Exact to Unresolved under the NOD pre-drill-else-main rule (decision 47)")
log("Table 4, row Exact", "7", "6: 15/9-19 A, F-10, F-15 A, F-15 B, F-15 C, F-15 D", "spud.exact", "15/9-F-11 moves to Unresolved under the NOD pre-drill-else-main rule (decision 47)")
log("Table 4, row Definitional offset", "(no wellbores listed)", "wellbores listed", "spud.definitional_offset", "detail added from the spud table")

W = nsdecls("w"); BORD = '<w:tcBorders><w:top w:val="single" w:color="808080" w:sz="4"/><w:left w:val="single" w:color="808080" w:sz="4"/><w:bottom w:val="single" w:color="808080" w:sz="4"/><w:right w:val="single" w:color="808080" w:sz="4"/></w:tcBorders>'
MAR = '<w:tcMar><w:top w:type="dxa" w:w="40"/><w:left w:type="dxa" w:w="80"/><w:bottom w:type="dxa" w:w="40"/><w:right w:type="dxa" w:w="80"/></w:tcMar>'


TBLPR0 = copy.deepcopy(doc.tables[2]._tbl.tblPr)      # table properties of the existing tables, kept before new tables are inserted


def make_table(header, rows, widths, hl=frozenset(), right=frozenset()):
    def cell(text, w, head=False, h=False, rt=False):
        shd = '<w:shd w:fill="E7E6E6" w:color="auto" w:val="clear"/>' if head else ""
        b = "<w:b/><w:bCs/>" if head else ""
        hh = '<w:highlight w:val="yellow"/>' if h else ""
        return (f'<w:tc><w:tcPr><w:tcW w:type="dxa" w:w="{w}"/>{BORD}{shd}{MAR}</w:tcPr><w:p><w:pPr><w:jc w:val="{"right" if rt else "left"}"/></w:pPr><w:r><w:rPr>'
                f'<w:rFonts w:ascii="Times New Roman" w:cs="Times New Roman" w:eastAsia="Times New Roman" w:hAnsi="Times New Roman"/>{b}<w:sz w:val="18"/><w:szCs w:val="18"/>{hh}</w:rPr>'
                f'<w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p></w:tc>')
    xml = "".join('<w:tr><w:trPr><w:cantSplit/><w:tblHeader/></w:trPr>' + "".join(cell(h, w, head=True) for h, w in zip(header, widths)) + "</w:tr>" for _ in [0])
    for ri, row in enumerate(rows):
        xml += '<w:tr><w:trPr><w:cantSplit/></w:trPr>' + "".join(cell(c, widths[ci], h=(ri, ci) in hl, rt=ci in right) for ci, c in enumerate(row)) + "</w:tr>"
    grid = "".join('<w:gridCol w:w="%d"/>' % w for w in widths)
    tbl = parse_xml(f'<w:tbl {W}><w:tblGrid>{grid}</w:tblGrid>{xml}</w:tbl>')
    tbl.insert(0, copy.deepcopy(TBLPR0))
    return tbl


fm = lambda x, d=0: f"{x:,.{d}f}"
def fdepth(x): return fm(x, 1) if float(x) % 1 else fm(x)
def fdiff(x):
    if x is None or (isinstance(x, float) and pd.isna(x)): return ""
    return "0" if float(x) == 0 else (f"{x:+,.1f}" if float(x) % 1 else f"{x:+,.0f}")

body_t, head_t, cap_t, note_t, bullet_t = P[48], P[53], P[55], P[56], P[98]


# ---- Abbreviations (final review) and reference entries (verified against the source PDFs, 28 October 2026)
def fp(prefix):
    hits = [p for p in P if p.text.startswith(prefix)]; assert len(hits) == 1, (prefix, len(hits)); return hits[0]


ABB = [("BHA", "bottom-hole assembly", "heading"), ("DP", "drill pipe", "DDR"), ("EZSV", "EZ Drill squeeze valve (cement retainer)", "DP"), ("LWD", "logging while drilling", "EZSV"),
       ("MT", "metric tonnes", "MD"), ("MWD", "measurement while drilling", "MT"), ("POOH", "pull out of hole", "P&A"), ("SG", "specific gravity", "RKB"),
       ("VSP", "vertical seismic profile", "TVD"), ("WOW", "waiting on weather", "WITSML")]
anchor_of = {"heading": [p for p in P if p.style.name.startswith("Heading") and p.text.strip() == "Abbreviations"][0]}
tmpl = fp("BOP  =")
for abbr, exp, after in ABB:
    anc = anchor_of[after] if after in anchor_of else fp(after + "  =")
    anchor_of[abbr] = new_par_after(anc._p, tmpl, f"{abbr}  =  {exp}", renum=False)
log("Abbreviations", "BOP, CI, DDR, MD, NOD, P&A, RKB, TD, TVD, WITSML", "adds BHA, DP, EZSV, LWD, MT, MWD, POOH, SG, VSP, WOW", None, "final review: abbreviations used in the quotations and in Table 2")


def ref_run_sub(prefix, old, new):
    p = fp(prefix)
    for r in p.runs:
        if old in r.text: r.text = r.text.replace(old, new); return
    raise AssertionError((prefix, old))


write_par(fp("Hoffimann, J."), "Hoffimann, J., Mao, Y., Wesley, A., and Taylor, A. 2018. Sequence Mining and Pattern Analysis in Drilling Reports with Deep Natural Language Processing. Presented at the SPE Annual Technical Conference and Exhibition, Dallas, Texas, USA, 24–26 September. SPE-191505-MS. https://doi.org/10.2118/191505-MS.")
ref_run_sub("Tunkiel, A.T.", "2020.", "2021.")
ref_run_sub("Priyadarshy, S.", "Houston, Texas, USA. OTC-27918-MS", "Houston, Texas, USA, 1–4 May. OTC-27918-MS")
ref_run_sub("Sidahmed, M.", "The Woodlands, Texas, USA. SPE-173429-MS", "The Woodlands, Texas, USA, 3–5 March. SPE-173429-MS")
log("References", "Hoffimann 2017 (arXiv); Tunkiel 2020; Priyadarshy and Sidahmed without conference days", "Hoffimann 2018 (SPE-191505-MS); Tunkiel 2021; conference days added (1-4 May; 3-5 March)", None, "all references verified against the source PDFs by the author on 28 October 2026 (decision 73)")

# ---- Table 4 rebuilt with a Rule column (the category rules moved out of the text of Section 5.5)
RULE = {"Exact agreement": "R4: WITSML date within 1 day of the NOD date, and the first report or free-text spud follows it by at most 1 day.",
        "Definitional offset": "R6: WITSML date within 8 days of the NOD date, and the free-text operational start follows it by 2 to 20 days. R5: a parent leg whose NOD record is claimed by a technical sidetrack and that has no free-text spud date of its own (15/9-19 B).",
        "WITSML data-entry error": "R2: WITSML date differs from the first report or free-text spud by more than 30 days, while the NOD date agrees with them within 3 days.",
        "NOD date marks later campaign": "R3: WITSML date within 7 days of the first report, but more than 1 year from the NOD date.",
        "Inherited parent spud date": "R1: the sidetrack carries the same spud date as its parent.",
        "Unresolved": "R5: a technical sidetrack whose NOD record is its parent leg's (15/9-19 BT2 with 19 B, 19 ST2 with 19 SR). R7: any wellbore that no other rule fits."}
rows4 = []
for r in t4.rows[1:]:
    c = [x.text for x in r.cells]
    if c[0] == "Wellbore-identity conflation": c[0] = "NOD date marks later campaign"
    rows4.append([c[0], c[1], RULE.get(c[0], ""), c[2]])
tbl4 = make_table(["Outcome", "Wellbores", "Rule", "Detail"], rows4, [1700, 800, 3900, 2960])
t4._tbl.addnext(tbl4); t4._tbl.getparent().remove(t4._tbl)
note4 = new_par_after(tbl4, note_t, "Rules are applied in the order R1 (inherited), R5 (technical sidetrack), R2 (error), R3 (later campaign), R4 (exact), R6 (definitional offset), R7 (unresolved). Unresolved includes 15/9-19 S, F-9 and F-11 (R7).")
log("Table 4", "three columns (Outcome, Wellbores, Detail); category named Wellbore-identity conflation", "four columns with a Rule column; category named NOD date marks later campaign", "spud.rule_thresholds_days;spud.table4",
    "TJ review: rules moved from the text of Section 5.5 into the table; the technical-sidetrack rule now names the wellbores it applies to (15/9-19 BT2, 19 ST2 and the parent leg 19 B)")

# ---- Table 1 (example report, Section 3.1) and Table 2 (phases, Section 3.3)
ex = val("example.ddr"); DISP = {"Reservoir drilling": "12¼ in. and smaller hole", "Surface casing": "Surface hole and casing"}
_full = " ".join(pd.read_parquet("data/interim/ddr.parquet", columns=["filename", "sum24Hr"]).set_index("filename").loc[ex["filename"], "sum24Hr"].split())
assert len(_full.split()) == ex["sum24Hr_words_total"] and all(f in _full for f in ex["sum24Hr_fragments"])          # the full summary, checked against the registry
sumtxt = "“" + _full + "”"; fctxt = " ... ".join(ex["forecast24hr_fragments"]) + " ..."
rows1 = [["wellbore", ex["wellbore"], "the wellbore the report belongs to"],
         ["date_report", ex["date_report"], "the day the report covers"],
         ["md (m)", fm(ex["md_m"]), "measured depth reached, along the hole"],
         ["dTimSpud", ex["dTimSpud"], "spud date in the report header, the same in every report of the wellbore"],
         ["md_kickoff", "absent", "kick-off depth field: empty in this report, as in most reports (Table 3)"],
         ["md_csg_last (m)", fdepth(ex["md_csg_last_m"]), "depth of the last casing shoe"],
         ["forecast24hr (abridged)", fctxt, "the plan for the next 24 hours"],
         ["sum24Hr", sumtxt, "free-text summary of the past 24 hours: the text the classifier reads"],
         ["Phase assigned", f"{DISP.get(ex['phase'], ex['phase'])} (priority {ex['phase_priority']})", "the classifier's result: the 17½ in. section matches the Intermediate drilling patterns, in either spelling of the size"]]
tbl1 = make_table(["Field", "Value", "Meaning"], rows1, [1700, 4300, 3360])
a1 = new_par_after(P[21]._p, body_t, "Table 1 shows one report as stored in the archive: the structured fields the pipeline uses and the free-text summary that the classifier reads, quoted in full. The example is a real construction report whose summary contains the one-half sign.", renum=False)
c1 = new_par_after(a1._p, cap_t, "Table 1: Example DDR as stored in the archive", renum=False)
c1._p.addnext(tbl1)
n1 = new_par_after(tbl1, note_t, f"Source: Equinor Volve dataset (Equinor 2018), file {ex['filename']}. The summary of {ex['sum24Hr_words_total']} words is quoted in full; the forecast is abridged. The kick-off field is empty in this report, as in most reports (Table 3).", renum=False)
log("Section 3.1", "(no example report)", "Table 1: one real report (15/9-F-5, 2007-12-21), summary quoted in full, with the fields explained", "example.ddr;table1.md_kickoff_m.pct", "TJ review: abridged example DDR")

GLYPH = {"1/2": "\u00bd", "1/4": "\u00bc", "3/4": "\u00be", "3/8": "\u215c", "5/8": "\u215d"}
READ = {"tagged sea bt": "tagged sea bottom"}           # the pattern matches the abbreviation seen in the report ("Tagged sea btm"); shown in full
rows2 = [[str(r["priority"]), DISP.get(r["phase"], r["phase"]).replace(" / ", "/"), ", ".join(re.sub(r"(\d+) (\d/\d)", lambda m: m.group(1) + GLYPH[m.group(2)], x) for x in r["sizes_in"]) or "none", ", ".join(READ.get(x, x) for x in r["words"]) or "(size only)"] for r in val("classifier.phase_table")]
tbl2 = make_table(["Priority", "Phase", "Hole and casing sizes (in.)", "Example trigger words"], rows2, [800, 2400, 2100, 4060])
a2 = new_par_after(P[25]._p, body_t, "Table 2 lists the 13 phases in priority order with representative triggers taken from the pattern table of the repository (outputs/stage4_priority_rules.csv). Sizes are matched in the spellings found in the reports, for example 17 1/2 or 17.5; words are matched without regard to case; each phase has further patterns that are not shown.", renum=False)
c2 = new_par_after(a2._p, cap_t, "Table 2: Phases in priority order with representative triggers", renum=False)
c2._p.addnext(tbl2)
n2 = new_par_after(tbl2, note_t, "A report is assigned to the highest-priority phase with a match; a report that matches no pattern is Unclassified.", renum=False)
log("Section 3.3", "(no phase table)", "Table 2: 13 phases, priority, sizes and trigger words from the frozen pattern table", "classifier.phase_table", "TJ review: generated from outputs/stage4_priority_rules.csv")

# ---- Table 5 (kick-off) and the text around it
kt = pd.DataFrame(val("kickoff.table"))
MET = {"depth in text": "depth in text", "event in text, depth from md": "event in text, depth from md", "depth sequence (narrative empty)": "depth sequence (narrative empty)"}
rows5 = []
for ri, r in enumerate(kt.itertuples()):
    parent = "15/9-19 S (per NOD)" if str(r.parent).startswith("none") else r.parent
    rows5.append([r.sidetrack, parent, fdepth(r.ddr_kickoff_depth_m), MET[r.method], "" if pd.isna(r.nod_wlbKickOffPoint_m) else fdepth(r.nod_wlbKickOffPoint_m), fdiff(r.nod_minus_ddr_m), r.confidence])
tbl5 = make_table(["Sidetrack", "Parent", "Kick-off depth (m MD)", "Method", "NOD kick-off (m)", "NOD minus DDR (m)", "Confidence"], rows5, [1350, 1500, 1050, 2260, 900, 1100, 1200], right={2, 4, 5})
A = new_par_after(P[67]._p, body_t, f"All 15 sidetracks by designation (15/9-19 A, 19 B, 19 BT2 and 19 ST2, F-1 A, F-1 B and F-1 C, F-9 A, F-11 A, F-11 B and F-11 T2, and F-15 A to F-15 D) yielded a kick-off depth (Table 5). For 11 the depth is written in the narrative; for 3 (15/9-F-1 A, F-1 B and F-15 D) the narrative names the kick-off event but not the depth, which is taken from the measured depth of the same report; for 1 (15/9-19 ST2) both reports at the transition are empty and the depth is the measured depth there, the weakest case. For 15/9-F-1 C the last F-1 B report of 22 February 2014 reads \"Commenced time drilling to1364 m. 50% formation obtained at 1364 m.\", and the first F-1 C report drills 17 1/2 in. hole from that depth, so the kick-off is 1,364 m. The depth of 3,059 m that the narratives report on 13 March 2014 is the start of the 8 1/2 in. section, after casing was run, not the kick-off. 15/9-19 A is a sidetrack whose kick-off is the top of a milled window at 2,178 m (window milled from 2,178 to 2,183 m), a depth that NOD also records.")
cap5 = new_par_after(A._p, cap_t, "Table 5: Sidetrack kick-off depths, 15 sidetracks")
cap5._p.addnext(tbl5)
note5 = new_par_after(tbl5, note_t, f"Depths are m MD. NOD minus DDR is NOD's wlbKickOffPoint minus the recovered depth; blank where NOD has no value. NOD lists 15/9-19 S as the parent; the re-entry from which 15/9-19 A was drilled (NOD 15/9-19 SR2, 1997) has no WITSML reports, so the kick-off is read from the 15/9-19 A reports.")
log("Section 5.6", "(no kick-off table)", "Table 5: sidetrack, parent, kick-off depth, method, NOD value, difference, confidence", "kickoff.table;kickoff.method_counts;kickoff.nod_comparison", "new table from the kick-off table (author-approved, no highlighting)")
log("Section 5.6", "Twelve free-text and two depth-sequence recoveries; 15th sidetrack unidentified; the F-1 C example 'Drilled and oriented 8 1/2\" from 3059 m'",
    "11 depth in text, 3 event in text with depth from md, 1 depth sequence; 15 recovered; F-1 C 1,364 m; 3,059 m is the 8 1/2 in. section start; 19 A (2,178 m) is the 15th sidetrack",
    "kickoff.method_counts;kickoff.proposed_depth.15/9-F-1 C;kickoff.proposed_depth.15/9-19 A", "decisions 27, 32, 40; open item 2 resolved")

# ---- 5.7 Duration
dt = pd.DataFrame(val("duration.table")); dl = pd.DataFrame(val("duration.drilled_length_and_rate_sidetracks")).set_index("wellbore"); do = pd.DataFrame(val("duration.drilled_length_and_rate_originals")).set_index("wellbore")
rows6 = []
for ri, r in enumerate(dt.itertuples()):
    src = dl if r.wellbore in dl.index else do
    ind = r.wellbore in val("duration.indicative_only")
    rows6.append([r.wellbore + (" (indicative)" if ind else ""), r.first_ddr, fdepth(r.td_md_m), r.td_date, fm(r.elapsed_days_to_td), fm(r.ddrs_to_td), "yes" if r.gap_flag else "", fdepth(src.loc[r.wellbore, "drilled_length_m"]), fm(src.loc[r.wellbore, "m_per_report_day"], 1)])
tbl6 = make_table(["Wellbore", "First report", "TD (m)", "TD reached", "Elapsed days to TD", "Reports to TD", "Gap flag", "Drilled length (m)", "m per report-day"], rows6, [1650, 1050, 800, 1050, 900, 800, 650, 1230, 1230], right={2, 4, 5, 7, 8})
h57 = new_par_after(P[69]._p, head_t, "5.7 Duration to Total Depth")
a57 = new_par_after(h57._p, body_t, f"Table 6 gives, for each wellbore, the first report (the 1980 date sentinel excluded), the date total depth was first reached (the maximum valid depth of Table 3), the elapsed calendar days from the first report to that date, the number of reports in that span (inclusive), the drilled length (total depth minus the kick-off depth of Table 5 for the 15 sidetracks, total depth for the other 11) and metres per report-day (drilled length divided by the number of reports). Elapsed days and reports agree where a wellbore was reported daily and differ where the archive has gaps; wellbores whose elapsed days exceed twice their reports are flagged. Seven of the 26 are flagged (15/9-19 S, F-1, F-4, F-5, F-9, F-14 and F-15). F-1 reaches its total depth of 3,632 m on 17 August 2013, 2,085 days after its first report on 2 December 2007, in 31 reports, four of them from 2007. 15/9-19 B (3,272 m) and 15/9-F-11 (347 m) end at the depth of an abandoned parent leg or stub (Section 5.4), so their durations are indicative only. Drilled lengths and rates for sidetracks depend on the kick-off depths of Section 5.6; rates are per reported day, not per elapsed day, and for F-1 they mix two campaigns. Per-phase report counts for each wellbore and well family are published with the duration table.")
cap6 = new_par_after(a57._p, cap_t, "Table 6: Duration to total depth (TD) and drilled length, 26 wellbores")
cap6._p.addnext(tbl6)
note6 = new_par_after(tbl6, note_t, "Gap flag: elapsed days to TD exceed twice the reports to TD. Drilled length is total depth minus the Table 5 kick-off depth for sidetracks and total depth for original wellbores; for the original wellbores this is measured depth from the rig datum and includes the water depth. (indicative): TD is the depth of an abandoned parent leg or stub.")
log("Section 5.7 (new)", "(no duration results)", "new subsection and Table 6: first report, TD, TD date, elapsed days, reports to TD, gap flag, drilled length, m per report-day", "duration.table;duration.drilled_length_and_rate_sidetracks;duration.drilled_length_and_rate_originals;duration.n_gap_flagged",
    "TJ answers to the Stage 9 questions")

# new conclusion bullet after the spud bullet (paragraph 98)
cb = new_par_after(P[98]._p, bullet_t, f"Sidetrack kick-off depths were recovered for all 15 sidetracks by one rule; for 15/9-F-1 C the kick-off is 1,364 m, not the 3,059 m start of the 8 1/2 in. section, and the four wellbores of the 15/9-F-1 family share their hole above that depth.")
log("Conclusions (new bullet)", "(none)", "kick-off result for all 15 sidetracks, F-1 C 1,364 m", "kickoff.proposed_depth.15/9-F-1 C;kickoff.F-1_family_shared_hole_above_m", "new result")

# draft banner after the affiliation note

# ---- Conclusions: summary paragraph before the numbered list (TJ review round 3, text exactly as supplied)
IH = [i for i in range(len(P)) if P[i].style.name.startswith("Heading") and P[i].text.strip() == "8 Conclusions"]; assert len(IH) == 1
concl = new_par_after(P[IH[0]]._p, body_t, "This paper shows three things. First, a field's daily drilling reports can be turned into a construction history that is checked against the regulator's register and whose error is measured, using a method anyone can inspect and rerun. Second, the structured fields of a WITSML archive cannot be trusted on their own: spud dates, kick-off depths and placeholder values all needed the narrative or the national register to correct them, and a kick-off read naively from the narrative can pick up the start of a hole section instead of the point where the sidetrack leaves its parent (3,059 m against 1,364 m for 15/9-F-1 C). Third, the national register alone does not give a complete inventory of what was drilled: abandoned parent legs, shared upper hole and missing campaigns appear only when the reports are reconstructed. For anyone screening old wells for CO2 storage or decommissioning, the practical lesson is to reconstruct the well history from the reports before relying on either source alone.", renum=False)
log("Conclusions", "(no summary paragraph)", "summary paragraph before the numbered list", "kickoff.proposed_depth.15/9-F-1 C", "TJ review round 3: text as supplied by the author")

# ---- Appendix B: audit arithmetic
hB = new_par_after(P[143]._p, P[142], "Appendix B: Audit Arithmetic", renum=False)
pB1 = new_par_after(hB._p, P[143], "This appendix gives the arithmetic behind the attainable classification rate of Section 5.2.", renum=False)
pB2 = new_par_after(pB1._p, P[143], POST(OLD50), renum=False)
log("Appendix B (new)", "(arithmetic was in Section 5.2)", "full arithmetic and stratum-weighted check", "audit.sample_proportion;audit.wilson_low;audit.wilson_high;audit.scale_frame_over_total;audit.residual_miss;audit.observed_rate;audit.stratum_weighted_rate_pct", "TJ review: moved from Section 5.2")

os.makedirs("manuscript", exist_ok=True)
OUT = os.environ.get("VOLVE_DRAFT_OUT", "manuscript/paper2_draft.docx")      # normally the tracked draft; another path when Word holds the file open
doc.save(OUT)

# =========================================================================== CHANGES.md and the checks
d2 = docx.Document(OUT)
hl_texts = []
for p in d2.paragraphs:
    for r in p.runs:
        if r.font.highlight_color is not None: hl_texts.append((p.text[:60], r.text))
for t in d2.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    if r.font.highlight_color is not None: hl_texts.append(("table", r.text))
open_notes = [t for _, t in hl_texts if t.startswith("[TJ:")]
prov = [t for _, t in hl_texts if not t.startswith("[TJ:")]
assert not prov, f"unexpected highlighted text: {prov}"      # no kick-off value carries highlighting after the author sign-off
json.dump(dict(log=LOG, resolved=RESOLVED, open_notes=open_notes, highlighted_non_note_runs=prov), open("data/interim/stage12_log.json", "w", encoding="utf-8"), indent=1, ensure_ascii=False)
print("saved", OUT, "|", len(LOG), "log entries |", len(RESOLVED), "resolved notes |", len(open_notes), "open notes |", len(prov), "highlighted values")
